"""Write result tables as CSV, styled XLSX, styled HTML and Markdown."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.evaluation.aggregate import best_mask
from src.utils.logger import get_logger

_log = get_logger(__name__)

HEADER_FILL = "1F3B4D"
BEST_FILL = "D7EFE9"
HTML_CSS = """
<style>
body{font-family:Inter,'Segoe UI',system-ui,sans-serif;margin:32px;color:#1d2733;background:#fbfbfa}
table{border-collapse:collapse;font-size:13px}
th{background:#1f3b4d;color:#fff;padding:8px 12px;text-align:left;font-weight:600}
td{padding:6px 12px;border-bottom:1px solid #e3e6ea;font-variant-numeric:tabular-nums}
tr:hover td{background:#f1f4f6}
td.best{background:#d7efe9;font-weight:700;color:#0e5e55}
</style>
"""


def write_all(tables: dict[str, pd.DataFrame], out_dir: Path) -> None:
    """Write every table to CSV, HTML and Markdown plus one multi-sheet XLSX workbook.

    Args:
        tables: Mapping of table name to DataFrame.
        out_dir: Destination folder.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, frame in tables.items():
        keep_index = name.startswith("pivot_")
        frame.to_csv(out_dir / f"{name}.csv", index=keep_index)
        (out_dir / f"{name}.md").write_text(
            frame.to_markdown(index=keep_index, floatfmt=".4f"), encoding="utf-8"
        )
        (out_dir / f"{name}.html").write_text(to_html(name, frame), encoding="utf-8")
    write_workbook(tables, out_dir / "results_tables.xlsx")
    _log.info("Wrote %d tables to %s (CSV, XLSX, HTML, Markdown)", len(tables), out_dir)


def to_html(name: str, frame: pd.DataFrame) -> str:
    """Render one table as a standalone styled HTML page.

    Args:
        name: Table name; pivot tables get their best value per column highlighted.
        frame: Table to render.

    Returns:
        A complete HTML document.
    """
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{name}</title>{HTML_CSS}" \
           f"</head><body><h2>{name.replace('_', ' ').title()}</h2>{table_fragment(name, frame)}" \
           "</body></html>"


def table_fragment(name: str, frame: pd.DataFrame) -> str:
    """Render a table as an HTML ``<table>`` fragment with best values marked.

    Args:
        name: Table name.
        frame: Table to render.

    Returns:
        An HTML fragment.
    """
    if not name.startswith("pivot_"):
        return frame.to_html(index=False, float_format=lambda v: f"{v:.4f}", border=0)
    mask = best_mask(frame, name.removeprefix("pivot_"))
    styler = frame.style.format("{:.4f}").set_td_classes(mask.map(lambda flag: "best" if flag else ""))
    return styler.to_html()


def write_workbook(tables: dict[str, pd.DataFrame], path: Path) -> None:
    """Write all tables to one workbook, one styled sheet per table.

    Args:
        tables: Mapping of table name to DataFrame.
        path: Destination ``.xlsx`` file.
    """
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in tables.items():
            sheet = name[:31]
            keep_index = name.startswith("pivot_")
            frame.to_excel(writer, sheet_name=sheet, index=keep_index)
            _style_sheet(writer.sheets[sheet], frame, name, keep_index)


def _style_sheet(sheet: object, frame: pd.DataFrame, name: str, has_index: bool) -> None:
    """Apply header colours, number formats, column widths and best-value fills.

    Args:
        sheet: openpyxl worksheet.
        frame: The table written to the sheet.
        name: Table name.
        has_index: Whether the first column is the DataFrame index.
    """
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center")
    offset = 2 if has_index else 1
    for column_index in range(1, sheet.max_column + 1):
        letter = get_column_letter(column_index)
        sheet.column_dimensions[letter].width = 18
        for cell in sheet[letter][1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    sheet.freeze_panes = "B2"
    if name.startswith("pivot_"):
        mask = best_mask(frame, name.removeprefix("pivot_"))
        for row_pos, column_pos in zip(*mask.to_numpy().nonzero()):
            cell = sheet.cell(row=row_pos + 2, column=column_pos + offset)
            cell.fill = PatternFill("solid", fgColor=BEST_FILL)
            cell.font = Font(bold=True)


def print_compact(master: pd.DataFrame) -> None:
    """Print the headline columns of the master table to the console.

    Args:
        master: Master results table.
    """
    columns = ["split", "model", "f1", "pr_auc", "mcc", "recall", "precision", "fit_seconds"]
    view = master[[c for c in columns if c in master.columns]]
    try:
        from rich.console import Console
        from rich.table import Table

        table = Table(title="Headline results (default threshold)", header_style="bold cyan")
        for column in view.columns:
            table.add_column(column, justify="right" if column not in ("split", "model") else "left")
        for row in view.itertuples(index=False):
            table.add_row(*[f"{v:.4f}" if isinstance(v, float) else str(v) for v in row])
        Console().print(table)
    except ImportError:
        print(view.to_markdown(index=False, floatfmt=".4f"))
